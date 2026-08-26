# -*- coding: utf-8 -*-
"""
Lager en statisk backup av /bergen-urban-decket som PPTX + PDF.

Kjører gjennom alle scenene i presentasjonen med Playwright, venter til hver
koreografi har landet, og fanger nøkkelframes i 3840x2160 (2x skala).
UI-kontroller (TOC-knapp, fullskjermknapp, fremdriftslinje) skjules i bildene.

Output:
  /app/public/digihome-presentasjon-backup.pptx
  /app/public/digihome-presentasjon-backup.pdf
"""
import os
import sys
import time

from playwright.sync_api import sync_playwright

BASE = "http://localhost:3000/bergen-urban"
FRAME_DIR = "/app/scripts/backup-frames"
PPTX_UT = "/app/public/digihome-presentasjon-backup.pptx"
PDF_UT = "/app/public/digihome-presentasjon-backup.pdf"

os.makedirs(FRAME_DIR, exist_ok=True)

# Hver oppføring: (filnavn, handlinger)
# Handlinger: ("toc", idx) = hopp via TOC-menyen · ("key", tast) = tastetrykk
#             ("wait", ms) = vent · ("shoot",) = ta bildet
SLIDES = [
    ("01-cover",            [("toc", 1),  ("wait", 5000),  ("shoot",)]),
    ("02-martin",           [("toc", 2),  ("wait", 6000),  ("shoot",)]),
    ("03-historien",        [("toc", 3),  ("wait", 7000),  ("shoot",)]),
    ("04-ideen",            [("toc", 4),  ("wait", 6000),  ("shoot",)]),
    ("05-prosessloopen",    [("toc", 5),  ("wait", 7000),  ("shoot",)]),
    ("06-problemet",        [("toc", 6),  ("wait", 12000), ("shoot",)]),
    ("07-losningen",        [("toc", 7),  ("wait", 6000),  ("shoot",)]),
    ("08-ai-utvikling",     [("toc", 8),  ("wait", 5500),  ("shoot",)]),
    # Prompten: steg 8 er tom — gå ett steg videre så teksten skrives ferdig
    ("09-prompten",         [("toc", 9),  ("wait", 2500), ("key", "ArrowRight"), ("wait", 6000), ("shoot",)]),
    # Agenten bygger: fang midt i kodeskrivingen (før deploy/auto-videre)
    ("10-agenten-bygger",   [("toc", 10), ("wait", 4500),  ("shoot",)]),
    # Portalen: tre beats — dashboard+assistent (0–8.6s), kalender (8.6–17.4s), enhet
    ("11-portal-dashboard", [("toc", 11), ("wait", 7500),  ("shoot",)]),
    ("12-portal-kalender",  [("wait", 6500), ("shoot",)]),
    ("13-portal-enhet",     [("wait", 8000), ("shoot",)]),
    ("14-produktveggen",    [("toc", 12), ("wait", 9000),  ("shoot",)]),
    # Integrasjonene: «alle»-fasen (alt lyser) inntreffer 13.1–19.1s
    ("15-integrasjonene",   [("toc", 13), ("wait", 15500), ("shoot",)]),
    ("16-kostnad-tittel",   [("toc", 14), ("wait", 5000),  ("shoot",)]),
    ("17-kostnad-tall",     [("key", "ArrowRight"), ("wait", 6000), ("shoot",)]),
    ("18-paastanden",       [("toc", 15), ("wait", 6000),  ("shoot",)]),
    ("19-seksaaringen",     [("toc", 16), ("wait", 6000),  ("shoot",)]),
    ("20-ai-agenter",       [("toc", 17), ("wait", 5500),  ("shoot",)]),
    ("21-ai-verktoy",       [("key", "ArrowRight"), ("wait", 6000), ("shoot",)]),
    ("22-book-mote",        [("toc", 19), ("wait", 7000),  ("shoot",)]),
]

SKJUL_JS = """
() => {
  const mode = window.__skjul ? 'hidden' : '';
  const fs = document.querySelector('button[title="Fullskjerm (F)"]');
  if (fs) fs.style.visibility = mode;
  const toc = document.querySelector('button[title="Innhold"]');
  if (toc && toc.parentElement) toc.parentElement.style.visibility = mode;
  document.querySelectorAll('div[aria-hidden="true"]').forEach((d) => {
    if (typeof d.className === 'string' && d.className.includes('bottom-0 z-40')) d.style.visibility = mode;
  });
}
"""


def toggle_chrome(page, skjul):
    page.evaluate(f"window.__skjul = {'true' if skjul else 'false'}")
    page.evaluate(SKJUL_JS)


def main():
    frames = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, executable_path="/usr/bin/chromium", args=["--no-sandbox"])
        ctx = browser.new_context(
            viewport={"width": 1920, "height": 1080},
            device_scale_factor=2,
        )
        page = ctx.new_page()
        print("Åpner decket ...", flush=True)
        page.goto(BASE, wait_until="networkidle", timeout=60000)
        page.wait_for_timeout(3000)

        for navn, handlinger in SLIDES:
            for h in handlinger:
                if h[0] == "toc":
                    page.click('button[title="Innhold"]', force=True)
                    page.wait_for_timeout(450)
                    page.click(f'[data-testid="bu-toc-item-{h[1]}"]', force=True)
                elif h[0] == "key":
                    page.keyboard.press(h[1])
                elif h[0] == "wait":
                    page.wait_for_timeout(h[1])
                elif h[0] == "shoot":
                    sti = os.path.join(FRAME_DIR, f"{navn}.jpg")
                    toggle_chrome(page, True)
                    page.wait_for_timeout(120)
                    page.screenshot(path=sti, type="jpeg", quality=92, full_page=False)
                    toggle_chrome(page, False)
                    frames.append(sti)
                    print(f"  fanget: {navn}", flush=True)
        browser.close()

    print(f"\n{len(frames)} frames fanget. Bygger PPTX ...", flush=True)

    # ── PPTX: 16:9, ett fullflate-bilde per slide ──
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for f in frames:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(f, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_UT)
    print(f"PPTX lagret: {PPTX_UT} ({os.path.getsize(PPTX_UT) // 1024} kB)", flush=True)

    # ── PDF: samme bilder, eksakte 16:9-sider ──
    import img2pdf

    side = (img2pdf.mm_to_pt(338.667), img2pdf.mm_to_pt(190.5))
    layout = img2pdf.get_layout_fun(side)
    with open(PDF_UT, "wb") as f:
        f.write(img2pdf.convert(frames, layout_fun=layout))
    print(f"PDF lagret: {PDF_UT} ({os.path.getsize(PDF_UT) // 1024} kB)", flush=True)
    print("FERDIG", flush=True)


if __name__ == "__main__":
    t0 = time.time()
    try:
        main()
    except Exception as e:
        print(f"FEIL: {e}", flush=True)
        sys.exit(1)
    print(f"Total tid: {time.time() - t0:.0f}s", flush=True)
