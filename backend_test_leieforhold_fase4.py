#!/usr/bin/env python3
"""
Backend test for LEIEFORHOLD FASE 1-4: FILTER/SCENARIO PÅ EKSPORT + FASTE KOSTNADER MED PERIODE
Tests export filters (xlsx/csv) and PUT /api/admin/leieforhold/okonomi/felles with startDato/sluttDato/aktiv fields.

Base URL: https://saker-hub.preview.emergentagent.com/api
Admin key: dh_admin_b3Kx92Qz7Lm4

CRITICAL SAFETY RULES:
(a) SendGrid is LIVE — do NOT trigger any email flows.
(b) Do NOT delete or modify the real cost entry 'Lønn — 1 ansatt' (belop 60000) in the enhetsokonomi collection.
    Create your own QA entries and delete ONLY those.
(c) Do NOT touch users qa-investor@example.com or martin@kviteberg.no.
(d) The platform API upstream is REAL PRODUCTION but all leieforhold calls are GET/read-only (safe).
(e) MongoDB: mongodb://localhost:27017 (db name from env). A previous run found a stale leieforhold_cache entry
    (source='contracts-fallback'); if GET /api/admin/leieforhold returns a non-'lease-income' source, clear the
    leieforhold_cache collection and retry with &fresh=1 before failing.
"""

import requests
import sys
from datetime import datetime, timedelta
from pymongo import MongoClient
import os
import zipfile
import io
from openpyxl import load_workbook

BASE_URL = "https://saker-hub.preview.emergentagent.com/api"
ADMIN_KEY = "dh_admin_b3Kx92Qz7Lm4"

# MongoDB connection
MONGO_URL = os.getenv("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.getenv("DB_NAME", "your_database_name")

def get_db():
    client = MongoClient(MONGO_URL)
    return client[DB_NAME]

def test_1_xlsx_unfiltered():
    """Test 1: GET /api/admin/leieforhold/xlsx without filters → 200, valid XLSX, no '-filtrert' in filename"""
    print("\n=== TEST 1: GET xlsx unfiltered ===")
    try:
        url = f"{BASE_URL}/admin/leieforhold/xlsx?key={ADMIN_KEY}"
        resp = requests.get(url, timeout=30)
        
        if resp.status_code != 200:
            print(f"❌ Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text[:500]}")
            return False
        
        # Check Content-Type
        content_type = resp.headers.get('Content-Type', '')
        if 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' not in content_type:
            print(f"❌ Expected xlsx content-type, got: {content_type}")
            return False
        
        # Check Content-Disposition does NOT contain '-filtrert'
        content_disp = resp.headers.get('Content-Disposition', '')
        if '-filtrert' in content_disp:
            print(f"❌ Unfiltered export should NOT have '-filtrert' in filename: {content_disp}")
            return False
        
        # Verify it's a valid XLSX (ZIP with PK signature)
        if not resp.content.startswith(b'PK'):
            print(f"❌ Response does not start with PK (ZIP signature)")
            return False
        
        # Load workbook with openpyxl (ignore warnings about extensions)
        try:
            wb = load_workbook(io.BytesIO(resp.content))
            sheet_names = wb.sheetnames
            
            # Verify required sheets exist
            required_sheets = ['Oversikt', 'Leieforhold', 'Per huseier', 'Økonomi']
            for sheet in required_sheets:
                if sheet not in sheet_names:
                    print(f"❌ Missing required sheet: {sheet}")
                    print(f"Found sheets: {sheet_names}")
                    return False
            
            # Verify 'Leieforhold' sheet has 15 columns with 'Utflytting' at column I (index 9)
            ws = wb['Leieforhold']
            # Header row is at row 9 according to the code
            header_row = 9
            headers = []
            for col in range(1, 16):  # A-O (15 columns)
                cell_value = ws.cell(row=header_row, column=col).value
                headers.append(cell_value)
            
            if len(headers) != 15:
                print(f"❌ Expected 15 header columns, got {len(headers)}")
                return False
            
            if 'Utflytting' not in headers:
                print(f"❌ 'Utflytting' column not found in headers: {headers}")
                return False
            
            # Verify 'Utflytting' is at column I (index 9, 0-based = 8)
            if headers[8] != 'Utflytting':
                print(f"❌ 'Utflytting' should be at column I (index 9), but found: {headers[8]}")
                return False
            
            print(f"✅ GET xlsx unfiltered: 200, valid XLSX, no '-filtrert', sheets: {sheet_names}")
            print(f"✅ 'Leieforhold' sheet has 15 columns with 'Utflytting' at column I")
            return True
            
        except Exception as e:
            print(f"❌ Failed to load workbook: {e}")
            return False
        
    except Exception as e:
        print(f"❌ Test 1 failed with exception: {e}")
        return False

def test_2_xlsx_filtered_status():
    """Test 2: GET xlsx with &status=leased,future → 200, '-filtrert' in filename, fewer rows"""
    print("\n=== TEST 2: GET xlsx with status filter ===")
    try:
        # First get unfiltered count
        url_unfiltered = f"{BASE_URL}/admin/leieforhold/xlsx?key={ADMIN_KEY}"
        resp_unfiltered = requests.get(url_unfiltered, timeout=30)
        if resp_unfiltered.status_code != 200:
            print(f"❌ Failed to get unfiltered xlsx for comparison")
            return False
        
        wb_unfiltered = load_workbook(io.BytesIO(resp_unfiltered.content))
        ws_unfiltered = wb_unfiltered['Leieforhold']
        # Count data rows (starting from row 10, first data row)
        unfiltered_rows = 0
        for row in range(10, ws_unfiltered.max_row + 1):
            if ws_unfiltered.cell(row=row, column=1).value:  # If first column has value
                unfiltered_rows += 1
        
        # Now get filtered
        url = f"{BASE_URL}/admin/leieforhold/xlsx?key={ADMIN_KEY}&status=leased,future"
        resp = requests.get(url, timeout=30)
        
        if resp.status_code != 200:
            print(f"❌ Expected 200, got {resp.status_code}")
            return False
        
        # Check Content-Disposition contains '-filtrert'
        content_disp = resp.headers.get('Content-Disposition', '')
        if '-filtrert' not in content_disp:
            print(f"❌ Filtered export should have '-filtrert' in filename: {content_disp}")
            return False
        
        # Verify it's valid XLSX
        if not resp.content.startswith(b'PK'):
            print(f"❌ Response does not start with PK (ZIP signature)")
            return False
        
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb['Leieforhold']
        
        # Count data rows
        filtered_rows = 0
        for row in range(10, ws.max_row + 1):
            if ws.cell(row=row, column=1).value:
                filtered_rows += 1
        
        if filtered_rows >= unfiltered_rows:
            print(f"❌ Filtered rows ({filtered_rows}) should be FEWER than unfiltered ({unfiltered_rows})")
            return False
        
        print(f"✅ GET xlsx with status=leased,future: 200, '-filtrert' in filename")
        print(f"✅ Filtered rows ({filtered_rows}) < unfiltered rows ({unfiltered_rows})")
        return True
        
    except Exception as e:
        print(f"❌ Test 2 failed with exception: {e}")
        return False

def test_3_xlsx_scenario():
    """Test 3: GET xlsx with &scenario=<future date> → 200, '-filtrert' in filename"""
    print("\n=== TEST 3: GET xlsx with scenario filter ===")
    try:
        # Use a date ~1 month in the future
        future_date = (datetime.now() + timedelta(days=30)).strftime('%Y-%m-%d')
        
        url = f"{BASE_URL}/admin/leieforhold/xlsx?key={ADMIN_KEY}&scenario={future_date}"
        resp = requests.get(url, timeout=30)
        
        if resp.status_code != 200:
            print(f"❌ Expected 200, got {resp.status_code}")
            return False
        
        # Check Content-Disposition contains '-filtrert'
        content_disp = resp.headers.get('Content-Disposition', '')
        if '-filtrert' not in content_disp:
            print(f"❌ Scenario export should have '-filtrert' in filename: {content_disp}")
            return False
        
        # Verify it's valid XLSX
        if not resp.content.startswith(b'PK'):
            print(f"❌ Response does not start with PK (ZIP signature)")
            return False
        
        print(f"✅ GET xlsx with scenario={future_date}: 200, '-filtrert' in filename")
        return True
        
    except Exception as e:
        print(f"❌ Test 3 failed with exception: {e}")
        return False

def test_4_csv_unfiltered_and_filtered():
    """Test 4: GET csv → 200, 15 columns including 'Utflytting'; GET csv with &status=leased → '-filtrert', fewer rows"""
    print("\n=== TEST 4: GET csv unfiltered and filtered ===")
    try:
        # First get unfiltered
        url_unfiltered = f"{BASE_URL}/admin/leieforhold/csv?key={ADMIN_KEY}"
        resp_unfiltered = requests.get(url_unfiltered, timeout=30)
        
        if resp_unfiltered.status_code != 200:
            print(f"❌ Expected 200 for unfiltered csv, got {resp_unfiltered.status_code}")
            return False
        
        # Check Content-Type
        content_type = resp_unfiltered.headers.get('Content-Type', '')
        if 'text/csv' not in content_type:
            print(f"❌ Expected text/csv content-type, got: {content_type}")
            return False
        
        # Parse CSV (UTF-8 with BOM)
        csv_text = resp_unfiltered.text
        if csv_text.startswith('\ufeff'):
            csv_text = csv_text[1:]  # Remove BOM
        
        lines = csv_text.strip().split('\n')
        if len(lines) < 1:
            print(f"❌ CSV has no lines")
            return False
        
        # Check header line (first line after BOM)
        header_line = lines[0]
        headers = header_line.split(';')
        
        if len(headers) != 15:
            print(f"❌ Expected 15 columns in CSV header, got {len(headers)}: {headers}")
            return False
        
        if 'Utflytting' not in headers:
            print(f"❌ 'Utflytting' column not found in CSV headers: {headers}")
            return False
        
        unfiltered_data_lines = len(lines) - 1  # Exclude header
        
        print(f"✅ GET csv unfiltered: 200, 15 columns including 'Utflytting', {unfiltered_data_lines} data lines")
        
        # Now get filtered
        url_filtered = f"{BASE_URL}/admin/leieforhold/csv?key={ADMIN_KEY}&status=leased"
        resp_filtered = requests.get(url_filtered, timeout=30)
        
        if resp_filtered.status_code != 200:
            print(f"❌ Expected 200 for filtered csv, got {resp_filtered.status_code}")
            return False
        
        # Check Content-Disposition contains '-filtrert'
        content_disp = resp_filtered.headers.get('Content-Disposition', '')
        if '-filtrert' not in content_disp:
            print(f"❌ Filtered CSV should have '-filtrert' in filename: {content_disp}")
            return False
        
        csv_text_filtered = resp_filtered.text
        if csv_text_filtered.startswith('\ufeff'):
            csv_text_filtered = csv_text_filtered[1:]
        
        lines_filtered = csv_text_filtered.strip().split('\n')
        filtered_data_lines = len(lines_filtered) - 1
        
        if filtered_data_lines >= unfiltered_data_lines:
            print(f"❌ Filtered data lines ({filtered_data_lines}) should be FEWER than unfiltered ({unfiltered_data_lines})")
            return False
        
        print(f"✅ GET csv with status=leased: 200, '-filtrert' in filename, {filtered_data_lines} data lines < {unfiltered_data_lines}")
        return True
        
    except Exception as e:
        print(f"❌ Test 4 failed with exception: {e}")
        return False

def test_5_put_felles_with_period():
    """Test 5: PUT /api/admin/leieforhold/okonomi/felles with startDato/sluttDato/aktiv fields"""
    print("\n=== TEST 5: PUT felles with period fields ===")
    qa_id = None
    try:
        db = get_db()
        
        # Create QA cost entry
        body = {
            "navn": "QA Periodekost",
            "belop": 5000,
            "fordeling": "alle",
            "kategori": "annet",
            "startDato": "2026-01-01",
            "sluttDato": "2026-12-31",
            "aktiv": True
        }
        
        url = f"{BASE_URL}/admin/leieforhold/okonomi/felles?key={ADMIN_KEY}"
        resp = requests.put(url, json=body, timeout=10)
        
        if resp.status_code != 200:
            print(f"❌ PUT felles expected 200, got {resp.status_code}: {resp.text}")
            return False
        
        data = resp.json()
        if not data.get('ok'):
            print(f"❌ PUT felles response not ok: {data}")
            return False
        
        qa_id = data.get('id')
        if not qa_id:
            print(f"❌ PUT felles did not return id: {data}")
            return False
        
        print(f"✅ PUT felles created QA cost entry with id: {qa_id}")
        
        # GET okonomi to verify fields
        url_get = f"{BASE_URL}/admin/leieforhold/okonomi?key={ADMIN_KEY}"
        resp_get = requests.get(url_get, timeout=10)
        
        if resp_get.status_code != 200:
            print(f"❌ GET okonomi expected 200, got {resp_get.status_code}")
            return False
        
        okonomi_data = resp_get.json()
        felles = okonomi_data.get('felles', [])
        
        qa_entry = None
        for entry in felles:
            if entry.get('id') == qa_id:
                qa_entry = entry
                break
        
        if not qa_entry:
            print(f"❌ QA entry not found in GET okonomi response")
            return False
        
        if qa_entry.get('startDato') != '2026-01-01':
            print(f"❌ startDato mismatch: expected '2026-01-01', got {qa_entry.get('startDato')}")
            return False
        
        if qa_entry.get('sluttDato') != '2026-12-31':
            print(f"❌ sluttDato mismatch: expected '2026-12-31', got {qa_entry.get('sluttDato')}")
            return False
        
        if qa_entry.get('aktiv') != True:
            print(f"❌ aktiv mismatch: expected True, got {qa_entry.get('aktiv')}")
            return False
        
        print(f"✅ GET okonomi shows QA entry with correct startDato, sluttDato, aktiv fields")
        
        # Update aktiv to false (need to provide navn and belop as they're required)
        body_update = {
            "id": qa_id,
            "navn": "QA Periodekost",
            "belop": 5000,
            "fordeling": "alle",
            "kategori": "annet",
            "aktiv": False
        }
        resp_update = requests.put(url, json=body_update, timeout=10)
        
        if resp_update.status_code != 200:
            print(f"❌ PUT aktiv=false expected 200, got {resp_update.status_code}")
            return False
        
        # GET again to verify
        resp_get2 = requests.get(url_get, timeout=10)
        okonomi_data2 = resp_get2.json()
        felles2 = okonomi_data2.get('felles', [])
        
        qa_entry2 = None
        for entry in felles2:
            if entry.get('id') == qa_id:
                qa_entry2 = entry
                break
        
        if not qa_entry2:
            print(f"❌ QA entry not found after update")
            return False
        
        if qa_entry2.get('aktiv') != False:
            print(f"❌ aktiv should be False after update, got {qa_entry2.get('aktiv')}")
            return False
        
        print(f"✅ PUT aktiv=false successful, GET confirms aktiv=false")
        
        # Test validation: sluttDato < startDato should return 400
        body_invalid = {
            "id": qa_id,
            "navn": "QA Periodekost",
            "belop": 5000,
            "fordeling": "alle",
            "kategori": "annet",
            "startDato": "2026-12-31",
            "sluttDato": "2026-01-01"
        }
        resp_invalid = requests.put(url, json=body_invalid, timeout=10)
        
        if resp_invalid.status_code != 400:
            print(f"❌ PUT with sluttDato < startDato should return 400, got {resp_invalid.status_code}")
            return False
        
        print(f"✅ PUT with sluttDato < startDato correctly returns 400")
        
        # DELETE the QA entry
        url_delete = f"{BASE_URL}/admin/leieforhold/okonomi/felles?key={ADMIN_KEY}&id={qa_id}"
        resp_delete = requests.delete(url_delete, timeout=10)
        
        if resp_delete.status_code != 200:
            print(f"❌ DELETE expected 200, got {resp_delete.status_code}")
            return False
        
        # GET to confirm deletion
        resp_get3 = requests.get(url_get, timeout=10)
        okonomi_data3 = resp_get3.json()
        felles3 = okonomi_data3.get('felles', [])
        
        for entry in felles3:
            if entry.get('id') == qa_id:
                print(f"❌ QA entry still exists after DELETE")
                return False
        
        print(f"✅ DELETE successful, GET confirms QA entry is gone")
        qa_id = None  # Mark as cleaned up
        
        return True
        
    except Exception as e:
        print(f"❌ Test 5 failed with exception: {e}")
        # Cleanup
        if qa_id:
            try:
                url_delete = f"{BASE_URL}/admin/leieforhold/okonomi/felles?key={ADMIN_KEY}&id={qa_id}"
                requests.delete(url_delete, timeout=10)
                print(f"Cleaned up QA entry {qa_id}")
            except:
                pass
        return False

def test_6_auth():
    """Test 6: GET xlsx and csv without key → 401"""
    print("\n=== TEST 6: Auth test ===")
    try:
        # Test xlsx without key
        url_xlsx = f"{BASE_URL}/admin/leieforhold/xlsx"
        resp_xlsx = requests.get(url_xlsx, timeout=10)
        
        if resp_xlsx.status_code != 401:
            print(f"❌ GET xlsx without key should return 401, got {resp_xlsx.status_code}")
            return False
        
        print(f"✅ GET xlsx without key returns 401")
        
        # Test csv without key
        url_csv = f"{BASE_URL}/admin/leieforhold/csv"
        resp_csv = requests.get(url_csv, timeout=10)
        
        if resp_csv.status_code != 401:
            print(f"❌ GET csv without key should return 401, got {resp_csv.status_code}")
            return False
        
        print(f"✅ GET csv without key returns 401")
        
        return True
        
    except Exception as e:
        print(f"❌ Test 6 failed with exception: {e}")
        return False

def test_7_regression():
    """Test 7: GET /api/admin/leieforhold → 200 {ok, source:'lease-income'}; GET okonomi → 'Lønn — 1 ansatt' intact"""
    print("\n=== TEST 7: Regression test ===")
    try:
        db = get_db()
        
        # GET /api/admin/leieforhold
        url = f"{BASE_URL}/admin/leieforhold?key={ADMIN_KEY}"
        resp = requests.get(url, timeout=30)
        
        if resp.status_code != 200:
            print(f"❌ GET leieforhold expected 200, got {resp.status_code}")
            return False
        
        data = resp.json()
        if not data.get('ok'):
            print(f"❌ GET leieforhold response not ok: {data}")
            return False
        
        source = data.get('source')
        
        # Check for stale cache (rule e)
        if source != 'lease-income':
            print(f"⚠️  Source is '{source}', not 'lease-income'. Clearing stale cache and retrying with &fresh=1...")
            
            # Clear leieforhold_cache collection
            db.leieforhold_cache.delete_many({})
            print(f"Cleared leieforhold_cache collection")
            
            # Retry with &fresh=1
            url_fresh = f"{BASE_URL}/admin/leieforhold?key={ADMIN_KEY}&fresh=1"
            resp_fresh = requests.get(url_fresh, timeout=30)
            
            if resp_fresh.status_code != 200:
                print(f"❌ GET leieforhold with &fresh=1 expected 200, got {resp_fresh.status_code}")
                return False
            
            data_fresh = resp_fresh.json()
            source = data_fresh.get('source')
            
            if source != 'lease-income':
                print(f"❌ After clearing cache and &fresh=1, source is still '{source}', not 'lease-income'")
                return False
            
            print(f"✅ After clearing cache, source is now 'lease-income'")
        else:
            print(f"✅ GET leieforhold returns 200 with source='lease-income'")
        
        # GET okonomi and verify 'Lønn — 1 ansatt' exists
        url_okonomi = f"{BASE_URL}/admin/leieforhold/okonomi?key={ADMIN_KEY}"
        resp_okonomi = requests.get(url_okonomi, timeout=10)
        
        if resp_okonomi.status_code != 200:
            print(f"❌ GET okonomi expected 200, got {resp_okonomi.status_code}")
            return False
        
        okonomi_data = resp_okonomi.json()
        felles = okonomi_data.get('felles', [])
        
        lonn_entry = None
        for entry in felles:
            if entry.get('navn') == 'Lønn — 1 ansatt' and entry.get('belop') == 60000:
                lonn_entry = entry
                break
        
        if not lonn_entry:
            print(f"❌ 'Lønn — 1 ansatt' (belop 60000) not found in okonomi felles")
            print(f"Found entries: {[e.get('navn') for e in felles]}")
            return False
        
        print(f"✅ GET okonomi returns 200 and 'Lønn — 1 ansatt' (belop 60000) is intact")
        
        return True
        
    except Exception as e:
        print(f"❌ Test 7 failed with exception: {e}")
        return False

def cleanup():
    """Cleanup: verify no QA cost entries remain and 'Lønn — 1 ansatt' is intact"""
    print("\n=== CLEANUP: Verify no QA entries remain ===")
    try:
        url = f"{BASE_URL}/admin/leieforhold/okonomi?key={ADMIN_KEY}"
        resp = requests.get(url, timeout=10)
        
        if resp.status_code != 200:
            print(f"⚠️  GET okonomi for cleanup check failed: {resp.status_code}")
            return False
        
        data = resp.json()
        felles = data.get('felles', [])
        
        qa_entries = [e for e in felles if e.get('navn', '').startswith('QA ')]
        
        if qa_entries:
            print(f"⚠️  Found {len(qa_entries)} QA entries remaining:")
            for entry in qa_entries:
                print(f"  - {entry.get('navn')} (id: {entry.get('id')})")
                # Try to delete
                try:
                    url_delete = f"{BASE_URL}/admin/leieforhold/okonomi/felles?key={ADMIN_KEY}&id={entry.get('id')}"
                    requests.delete(url_delete, timeout=10)
                    print(f"  Deleted {entry.get('navn')}")
                except Exception as e:
                    print(f"  Failed to delete: {e}")
        else:
            print(f"✅ No QA entries found")
        
        # Verify 'Lønn — 1 ansatt' is intact
        lonn_entry = None
        for entry in felles:
            if entry.get('navn') == 'Lønn — 1 ansatt' and entry.get('belop') == 60000:
                lonn_entry = entry
                break
        
        if not lonn_entry:
            print(f"❌ CRITICAL: 'Lønn — 1 ansatt' (belop 60000) is MISSING!")
            return False
        
        print(f"✅ 'Lønn — 1 ansatt' (belop 60000) is intact")
        
        return True
        
    except Exception as e:
        print(f"⚠️  Cleanup verification failed: {e}")
        return False

def main():
    print("=" * 80)
    print("BACKEND TEST: LEIEFORHOLD FASE 1-4 - FILTER/SCENARIO + FASTE KOSTNADER MED PERIODE")
    print("=" * 80)
    print(f"Base URL: {BASE_URL}")
    print(f"Admin key: {ADMIN_KEY}")
    print(f"MongoDB: {MONGO_URL}, DB: {DB_NAME}")
    print("=" * 80)
    
    results = []
    
    # Run tests
    results.append(("Test 1: GET xlsx unfiltered", test_1_xlsx_unfiltered()))
    results.append(("Test 2: GET xlsx with status filter", test_2_xlsx_filtered_status()))
    results.append(("Test 3: GET xlsx with scenario", test_3_xlsx_scenario()))
    results.append(("Test 4: GET csv unfiltered and filtered", test_4_csv_unfiltered_and_filtered()))
    results.append(("Test 5: PUT felles with period fields", test_5_put_felles_with_period()))
    results.append(("Test 6: Auth test", test_6_auth()))
    results.append(("Test 7: Regression test", test_7_regression()))
    
    # Cleanup
    cleanup()
    
    # Summary
    print("\n" + "=" * 80)
    print("TEST SUMMARY")
    print("=" * 80)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for test_name, result in results:
        status = "✅ PASS" if result else "❌ FAIL"
        print(f"{status}: {test_name}")
    
    print("=" * 80)
    print(f"TOTAL: {passed}/{total} tests passed ({int(passed/total*100)}% success rate)")
    print("=" * 80)
    
    if passed == total:
        print("\n🎉 ALL TESTS PASSED!")
        sys.exit(0)
    else:
        print(f"\n⚠️  {total - passed} test(s) failed")
        sys.exit(1)

if __name__ == "__main__":
    main()
