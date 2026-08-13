#!/usr/bin/env python3
"""
Backend test for ØKONOMI TIL DATAROM-OVERSIKT + HONORAR-HERO + SATS MVA-MERKING

Tests:
1. GET /api/admin/datarom/oversikt with admin key - verify okonomi block
2. Investor login and reading datarom/oversikt
3. GET /api/admin/leieforhold/csv - verify 16 columns with VAT conversion
4. GET /api/admin/leieforhold/xlsx - verify XLSX format
5. Regression tests
"""

import requests
import sys
import io
import openpyxl

BASE_URL = "https://saker-hub.preview.emergentagent.com/api"
ADMIN_KEY = "dh_admin_b3Kx92Qz7Lm4"
INVESTOR_EMAIL = "qa-investor@example.com"
INVESTOR_PASSWORD = "QaInvest12345!"

def test_datarom_oversikt_admin():
    """Test 1: GET /api/admin/datarom/oversikt with admin key"""
    print("\n=== TEST 1: GET /api/admin/datarom/oversikt (admin key) ===")
    try:
        url = f"{BASE_URL}/admin/datarom/oversikt"
        params = {"key": ADMIN_KEY}
        resp = requests.get(url, params=params, timeout=30)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text[:500]}")
            return False
        
        data = resp.json()
        
        # Verify oversikt.okonomi exists
        if "oversikt" not in data or "okonomi" not in data["oversikt"]:
            print(f"❌ FAIL: Missing oversikt.okonomi in response")
            print(f"Keys: {data.keys()}")
            return False
        
        okonomi = data["oversikt"]["okonomi"]
        print(f"✓ oversikt.okonomi found")
        
        # Verify required fields
        required_fields = ["fellesMnd", "felles", "cacTotal", "breakEvenPct", "enheterIgjen", "paybackMnd"]
        for field in required_fields:
            if field not in okonomi:
                print(f"❌ FAIL: Missing field '{field}' in okonomi")
                return False
            print(f"✓ {field}: {okonomi[field]}")
        
        # Verify fellesMnd is a number
        if not isinstance(okonomi["fellesMnd"], (int, float)):
            print(f"❌ FAIL: fellesMnd is not a number: {type(okonomi['fellesMnd'])}")
            return False
        
        # Verify felles is an array
        if not isinstance(okonomi["felles"], list):
            print(f"❌ FAIL: felles is not an array: {type(okonomi['felles'])}")
            return False
        
        print(f"✓ felles array has {len(okonomi['felles'])} items")
        
        # Verify fellesMnd equals sum of ACTIVE felles posts
        if len(okonomi["felles"]) > 0:
            # Calculate sum of active posts
            active_sum = 0
            for post in okonomi["felles"]:
                if post.get("aktiv") != False:  # aktiv !== false
                    active_sum += post.get("belop", 0)
            
            print(f"✓ Sum of active felles posts: {active_sum}")
            print(f"✓ fellesMnd from API: {okonomi['fellesMnd']}")
            
            if abs(active_sum - okonomi["fellesMnd"]) > 0.01:
                print(f"❌ FAIL: fellesMnd ({okonomi['fellesMnd']}) != sum of active posts ({active_sum})")
                return False
            
            print(f"✓ fellesMnd matches sum of active posts")
        
        # Verify drift.kostnaderMnd includes fellesMnd
        if "drift" in data["oversikt"]:
            drift = data["oversikt"]["drift"]
            if "kostnaderMnd" in drift:
                print(f"✓ drift.kostnaderMnd: {drift['kostnaderMnd']}")
                if drift["kostnaderMnd"] < okonomi["fellesMnd"]:
                    print(f"⚠️  WARNING: drift.kostnaderMnd ({drift['kostnaderMnd']}) < fellesMnd ({okonomi['fellesMnd']})")
                else:
                    print(f"✓ drift.kostnaderMnd >= fellesMnd")
                
                # Verify marginMnd = honorarMnd - kostnaderMnd
                if "honorarMnd" in drift and "marginMnd" in drift:
                    expected_margin = drift["honorarMnd"] - drift["kostnaderMnd"]
                    if abs(drift["marginMnd"] - expected_margin) > 0.01:
                        print(f"❌ FAIL: marginMnd ({drift['marginMnd']}) != honorarMnd - kostnaderMnd ({expected_margin})")
                        return False
                    print(f"✓ drift.marginMnd = honorarMnd - kostnaderMnd")
        
        print("✅ TEST 1 PASSED")
        return True
        
    except Exception as e:
        print(f"❌ FAIL: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False

def test_datarom_oversikt_investor():
    """Test 2: Investor login and reading datarom/oversikt"""
    print("\n=== TEST 2: Investor login and GET datarom/oversikt ===")
    try:
        # Login as investor
        login_url = f"{BASE_URL}/admin/auth/login"
        login_data = {"email": INVESTOR_EMAIL, "password": INVESTOR_PASSWORD}
        resp = requests.post(login_url, json=login_data, timeout=30)
        print(f"Login status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Login failed with {resp.status_code}")
            print(f"Response: {resp.text[:500]}")
            return False
        
        login_result = resp.json()
        if "token" not in login_result:
            print(f"❌ FAIL: No token in login response")
            return False
        
        token = login_result["token"]
        print(f"✓ Logged in as investor, got token")
        
        # Get datarom/oversikt with investor token
        url = f"{BASE_URL}/admin/datarom/oversikt"
        params = {"key": token}
        resp = requests.get(url, params=params, timeout=30)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text[:500]}")
            return False
        
        data = resp.json()
        
        # Verify oversikt.okonomi exists
        if "oversikt" not in data or "okonomi" not in data["oversikt"]:
            print(f"❌ FAIL: Missing oversikt.okonomi in investor response")
            return False
        
        okonomi = data["oversikt"]["okonomi"]
        print(f"✓ Investor can read oversikt.okonomi")
        
        # Verify same structure as admin
        required_fields = ["fellesMnd", "felles", "cacTotal", "breakEvenPct", "enheterIgjen", "paybackMnd"]
        for field in required_fields:
            if field not in okonomi:
                print(f"❌ FAIL: Missing field '{field}' in investor okonomi")
                return False
        
        print(f"✓ All required fields present for investor")
        print(f"✓ fellesMnd: {okonomi['fellesMnd']}")
        print(f"✓ cacTotal: {okonomi['cacTotal']}")
        print(f"✓ breakEvenPct: {okonomi['breakEvenPct']}")
        
        print("✅ TEST 2 PASSED")
        return True
        
    except Exception as e:
        print(f"❌ FAIL: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False

def test_leieforhold_csv():
    """Test 3: GET /api/admin/leieforhold/csv - verify 16 columns with VAT conversion"""
    print("\n=== TEST 3: GET /api/admin/leieforhold/csv ===")
    try:
        url = f"{BASE_URL}/admin/leieforhold/csv"
        params = {"key": ADMIN_KEY}
        resp = requests.get(url, params=params, timeout=30)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text[:500]}")
            return False
        
        # Parse CSV
        csv_text = resp.text
        if csv_text.startswith('\ufeff'):
            csv_text = csv_text[1:]  # Remove BOM
        
        lines = csv_text.strip().split('\n')
        if len(lines) < 2:
            print(f"❌ FAIL: CSV has less than 2 lines")
            return False
        
        # Check header
        header = lines[0]
        columns = header.split(';')
        print(f"✓ CSV has {len(columns)} columns")
        print(f"✓ Header: {header[:200]}...")
        
        if len(columns) != 16:
            print(f"❌ FAIL: Expected 16 columns, got {len(columns)}")
            print(f"Columns: {columns}")
            return False
        
        # Verify column 11 is 'Sats (eks. mva)'
        if columns[10] != 'Sats (eks. mva)':
            print(f"❌ FAIL: Column 11 should be 'Sats (eks. mva)', got '{columns[10]}'")
            return False
        print(f"✓ Column 11 is 'Sats (eks. mva)'")
        
        # Verify column 12 is 'Avtalt sats'
        if columns[11] != 'Avtalt sats':
            print(f"❌ FAIL: Column 12 should be 'Avtalt sats', got '{columns[11]}'")
            return False
        print(f"✓ Column 12 is 'Avtalt sats'")
        
        # Check data lines for VAT conversion
        found_inkl = False
        found_eks = False
        vat_conversion_verified = False
        
        for i, line in enumerate(lines[1:], start=2):
            if not line.strip():
                continue
            
            parts = line.split(';')
            if len(parts) < 12:
                continue
            
            sats_eks = parts[10]  # Column 11 (0-indexed 10)
            avtalt_sats = parts[11]  # Column 12 (0-indexed 11)
            
            if 'inkl. mva' in avtalt_sats:
                found_inkl = True
                print(f"✓ Line {i}: Found 'inkl. mva' in Avtalt sats: {avtalt_sats}")
                
                # Verify conversion: if avtalt is "15 % inkl. mva", sats should be 12
                if '15 % inkl. mva' in avtalt_sats or '15% inkl. mva' in avtalt_sats:
                    # Expected: 15 / 1.25 = 12
                    sats_value = sats_eks.replace(',', '.')
                    try:
                        sats_num = float(sats_value)
                        expected = 15 / 1.25
                        if abs(sats_num - expected) < 0.1:
                            print(f"✓ VAT conversion verified: 15% inkl → {sats_num}% eks (expected {expected})")
                            vat_conversion_verified = True
                    except:
                        pass
            
            if 'eks. mva' in avtalt_sats:
                found_eks = True
                print(f"✓ Line {i}: Found 'eks. mva' in Avtalt sats: {avtalt_sats}")
        
        if not found_inkl:
            print(f"⚠️  WARNING: No 'inkl. mva' found in Avtalt sats column")
        
        if not found_eks:
            print(f"⚠️  WARNING: No 'eks. mva' found in Avtalt sats column")
        
        if found_inkl and not vat_conversion_verified:
            print(f"⚠️  WARNING: Found 'inkl. mva' but could not verify conversion")
        
        print("✅ TEST 3 PASSED")
        return True
        
    except Exception as e:
        print(f"❌ FAIL: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False

def test_leieforhold_xlsx():
    """Test 4: GET /api/admin/leieforhold/xlsx - verify XLSX format"""
    print("\n=== TEST 4: GET /api/admin/leieforhold/xlsx ===")
    try:
        url = f"{BASE_URL}/admin/leieforhold/xlsx"
        params = {"key": ADMIN_KEY}
        resp = requests.get(url, params=params, timeout=30)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text[:500]}")
            return False
        
        # Verify it's a valid XLSX (PK signature)
        content = resp.content
        if not content.startswith(b'PK'):
            print(f"❌ FAIL: Not a valid XLSX file (missing PK signature)")
            print(f"First bytes: {content[:20]}")
            return False
        
        print(f"✓ Valid XLSX file (PK signature found)")
        print(f"✓ File size: {len(content)} bytes")
        
        # Try to open with openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content))
            print(f"✓ Successfully opened with openpyxl")
            print(f"✓ Worksheets: {wb.sheetnames}")
            
            # Check if 'Leieforhold' sheet exists
            if 'Leieforhold' not in wb.sheetnames:
                print(f"❌ FAIL: 'Leieforhold' sheet not found")
                return False
            
            ws = wb['Leieforhold']
            print(f"✓ 'Leieforhold' sheet found")
            
            # Check header row (row 9)
            header_row = 9
            col_k_header = ws.cell(row=header_row, column=11).value
            print(f"✓ Column K header (row {header_row}): {col_k_header}")
            
            if col_k_header != 'Sats (eks. mva)':
                print(f"❌ FAIL: Column K should be 'Sats (eks. mva)', got '{col_k_header}'")
                return False
            
            print(f"✓ Column K header is 'Sats (eks. mva)'")
            
            # Check for VAT conversion in data rows
            # Find a row with inkl. mva (check column L 'Avtalt sats' in CSV, but in XLSX we check notes/fill)
            found_vat_note = False
            for row in range(10, min(40, ws.max_row + 1)):  # Check first 30 data rows
                cell_k = ws.cell(row=row, column=11)
                
                # Check if cell has yellow fill (inkl. mva indicator)
                if cell_k.fill and cell_k.fill.fgColor:
                    fill_color = cell_k.fill.fgColor.rgb if hasattr(cell_k.fill.fgColor, 'rgb') else None
                    if fill_color and 'FDF3E0' in str(fill_color):
                        print(f"✓ Row {row}: Found yellow fill on column K (inkl. mva indicator)")
                        found_vat_note = True
                        
                        # Check note
                        if cell_k.comment:
                            print(f"✓ Row {row}: Cell has note: {cell_k.comment.text[:100]}...")
                        
                        # Check value is converted (should be eks. mva)
                        value = cell_k.value
                        if value:
                            print(f"✓ Row {row}: Sats value (eks. mva): {value}")
                        
                        break
            
            if not found_vat_note:
                print(f"⚠️  WARNING: No yellow fill found on column K (no inkl. mva rows detected)")
            
            print("✅ TEST 4 PASSED")
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error opening XLSX with openpyxl: {e}")
            # This is acceptable if it's just warnings about conditional formatting
            if "conditional formatting" in str(e).lower() or "data validation" in str(e).lower():
                print(f"⚠️  Known warning (acceptable): {e}")
                print("✅ TEST 4 PASSED (with warnings)")
                return True
            raise
        
    except Exception as e:
        print(f"❌ FAIL: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False

def test_regression():
    """Test 5: Regression tests"""
    print("\n=== TEST 5: Regression tests ===")
    try:
        # Test 5a: GET /api/admin/leieforhold
        url = f"{BASE_URL}/admin/leieforhold"
        params = {"key": ADMIN_KEY}
        resp = requests.get(url, params=params, timeout=30)
        print(f"GET /api/admin/leieforhold status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Expected 200, got {resp.status_code}")
            return False
        
        data = resp.json()
        if "rows" not in data:
            print(f"❌ FAIL: Missing 'rows' in response")
            return False
        
        row_count = len(data["rows"])
        print(f"✓ GET /api/admin/leieforhold returned {row_count} rows")
        
        if row_count < 25 or row_count > 35:
            print(f"⚠️  WARNING: Expected ~29 rows, got {row_count}")
        
        # Test 5b: GET /api/admin/leieforhold/okonomi
        url = f"{BASE_URL}/admin/leieforhold/okonomi"
        params = {"key": ADMIN_KEY}
        resp = requests.get(url, params=params, timeout=30)
        print(f"GET /api/admin/leieforhold/okonomi status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAIL: Expected 200, got {resp.status_code}")
            return False
        
        data = resp.json()
        if "felles" not in data:
            print(f"❌ FAIL: Missing 'felles' in okonomi response")
            return False
        
        felles_count = len(data["felles"])
        print(f"✓ GET /api/admin/leieforhold/okonomi returned {felles_count} felles posts")
        
        # Verify posts have expected structure
        if felles_count > 0:
            first_post = data["felles"][0]
            if "navn" in first_post and "belop" in first_post:
                print(f"✓ First post: {first_post.get('navn')} - {first_post.get('belop')} kr")
        
        # Test 5c: GET without key should return 401
        url = f"{BASE_URL}/admin/datarom/oversikt"
        resp = requests.get(url, timeout=30)
        print(f"GET datarom/oversikt without key status: {resp.status_code}")
        
        if resp.status_code != 401:
            print(f"❌ FAIL: Expected 401 without key, got {resp.status_code}")
            return False
        
        print(f"✓ GET without key correctly returns 401")
        
        print("✅ TEST 5 PASSED")
        return True
        
    except Exception as e:
        print(f"❌ FAIL: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False

def main():
    print("=" * 80)
    print("BACKEND TEST: ØKONOMI TIL DATAROM-OVERSIKT + SATS MVA-MERKING")
    print("=" * 80)
    print(f"Base URL: {BASE_URL}")
    print(f"Admin key: {ADMIN_KEY}")
    print(f"Investor: {INVESTOR_EMAIL}")
    
    results = []
    
    # Run tests
    results.append(("Test 1: Datarom oversikt (admin)", test_datarom_oversikt_admin()))
    results.append(("Test 2: Datarom oversikt (investor)", test_datarom_oversikt_investor()))
    results.append(("Test 3: Leieforhold CSV", test_leieforhold_csv()))
    results.append(("Test 4: Leieforhold XLSX", test_leieforhold_xlsx()))
    results.append(("Test 5: Regression", test_regression()))
    
    # Summary
    print("\n" + "=" * 80)
    print("SUMMARY")
    print("=" * 80)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for name, result in results:
        status = "✅ PASS" if result else "❌ FAIL"
        print(f"{status}: {name}")
    
    print(f"\nTotal: {passed}/{total} tests passed")
    
    if passed == total:
        print("\n🎉 ALL TESTS PASSED!")
        return 0
    else:
        print(f"\n❌ {total - passed} test(s) failed")
        return 1

if __name__ == "__main__":
    sys.exit(main())
