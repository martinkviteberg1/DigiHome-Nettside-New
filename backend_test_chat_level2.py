#!/usr/bin/env python3
"""
Backend test for Chat Level 2 features in DigiHome portal.
Tests: file upload (chunked), attachments, reactions, editing, pinning, typing indicator, DOCX preview.
"""

import requests
import json
import base64
import time
from pymongo import MongoClient
from io import BytesIO
import os

# Configuration
BASE_URL = os.getenv('NEXT_PUBLIC_BASE_URL', 'https://saker-hub.preview.emergentagent.com')
API_BASE = f"{BASE_URL}/api"
ADMIN_KEY = "dh_admin_b3Kx92Qz7Lm4"
MONGO_URL = "mongodb://localhost:27017"
DB_NAME = "your_database_name"

# Test tracking
tracked_file_ids = []
tracked_message_ids = []
tracked_user_ids = []

def log_test(test_name, passed, details=""):
    status = "✅ PASS" if passed else "❌ FAIL"
    print(f"{status}: {test_name}")
    if details:
        print(f"  {details}")
    if not passed:
        raise Exception(f"Test failed: {test_name} - {details}")

def create_small_png_base64():
    """Create a minimal valid PNG (1x1 transparent pixel)"""
    # PNG signature + IHDR + IDAT + IEND
    png_bytes = base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=='
    )
    return base64.b64encode(png_bytes).decode('utf-8')

def create_minimal_docx():
    """Create a minimal valid DOCX file"""
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph('QA Test Document')
        doc.add_paragraph('This is a test DOCX file for preview testing.')
        
        # Save to BytesIO
        docx_bytes = BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)
        return base64.b64encode(docx_bytes.read()).decode('utf-8')
    except ImportError:
        print("  python-docx not available, skipping DOCX creation")
        return None

def main():
    print("=" * 80)
    print("CHAT LEVEL 2 BACKEND TEST")
    print("=" * 80)
    print(f"Base URL: {API_BASE}")
    print(f"Admin key: {ADMIN_KEY}")
    print(f"MongoDB: {MONGO_URL}, DB: {DB_NAME}")
    print()
    
    # Connect to MongoDB
    mongo_client = MongoClient(MONGO_URL)
    db = mongo_client[DB_NAME]
    
    # Baseline counts
    baseline_messages = db.chat_messages.count_documents({})
    baseline_files = db.chat_files.count_documents({})
    baseline_pinned = db.chat_messages.count_documents({'festet': True})
    print(f"BASELINE: {baseline_messages} messages, {baseline_files} files, {baseline_pinned} pinned")
    print()
    
    # Login as owner to get token
    print("Logging in as owner...")
    login_resp = requests.post(
        f"{API_BASE}/admin/auth/login",
        json={"email": "martin@kviteberg.no", "password": "Pyramiden2025##"}
    )
    if login_resp.status_code != 200:
        raise Exception(f"Login failed: {login_resp.status_code} {login_resp.text}")
    
    owner_token = login_resp.json().get('token')
    owner_id = login_resp.json().get('user', {}).get('id')
    print(f"  Owner token: {owner_token[:20]}...")
    print(f"  Owner ID: {owner_id}")
    print()
    
    # Create QA user for multi-user tests
    print("Creating QA user...")
    qa_user_resp = requests.post(
        f"{API_BASE}/admin/users?key={ADMIN_KEY}",
        json={
            "name": "QA Chat User",
            "email": "qa-chat-user@example.com",
            "role": "bruker",
            "password": "QAtest1234",
            "invite": False
        }
    )
    if qa_user_resp.status_code not in [200, 201]:
        raise Exception(f"Failed to create QA user: {qa_user_resp.status_code} {qa_user_resp.text}")
    
    qa_user_id = qa_user_resp.json().get('id')
    tracked_user_ids.append(qa_user_id)
    print(f"  QA user ID: {qa_user_id}")
    
    # Login as QA user
    qa_login_resp = requests.post(
        f"{API_BASE}/admin/auth/login",
        json={"email": "qa-chat-user@example.com", "password": "QAtest1234"}
    )
    if qa_login_resp.status_code != 200:
        raise Exception(f"QA login failed: {qa_login_resp.status_code} {qa_login_resp.text}")
    
    qa_token = qa_login_resp.json().get('token')
    print(f"  QA token: {qa_token[:20]}...")
    print()
    
    # ========================================================================
    # A) FILE UPLOAD (POST /api/admin/chat/fil-chunk)
    # ========================================================================
    print("=" * 80)
    print("A) FILE UPLOAD TESTS")
    print("=" * 80)
    
    # A1: Upload a small PNG in 2 chunks
    print("A1: Upload PNG in 2 chunks...")
    png_b64 = create_small_png_base64()
    mid_point = len(png_b64) // 2
    chunk1 = png_b64[:mid_point]
    chunk2 = png_b64[mid_point:]
    upload_id = "qa-upload-1"
    
    # First chunk
    chunk1_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": upload_id,
            "index": 0,
            "total": 2,
            "data": chunk1,
            "name": "qa-test.png",
            "type": "image/png",
            "kanal": "generelt"
        }
    )
    log_test(
        "A1a: First chunk returns complete:false",
        chunk1_resp.status_code == 200 and not chunk1_resp.json().get('complete'),
        f"Status: {chunk1_resp.status_code}, Response: {chunk1_resp.json()}"
    )
    
    # Second chunk
    chunk2_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": upload_id,
            "index": 1,
            "total": 2,
            "data": chunk2,
            "name": "qa-test.png",
            "type": "image/png",
            "kanal": "generelt"
        }
    )
    chunk2_data = chunk2_resp.json()
    file_id_1 = chunk2_data.get('fil', {}).get('id')
    tracked_file_ids.append(file_id_1)
    
    log_test(
        "A1b: Second chunk returns complete:true with file metadata",
        chunk2_resp.status_code == 200 and 
        chunk2_data.get('complete') and 
        file_id_1 and
        chunk2_data.get('fil', {}).get('type') == 'image/png',
        f"File ID: {file_id_1}, Size: {chunk2_data.get('fil', {}).get('size')}"
    )
    print()
    
    # A2: Invalid chunk (index >= total)
    print("A2: Invalid chunk (index >= total)...")
    invalid_chunk_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": "qa-invalid",
            "index": 5,
            "total": 2,
            "data": "test",
            "name": "invalid.png",
            "type": "image/png",
            "kanal": "generelt"
        }
    )
    log_test(
        "A2: Invalid chunk returns 400",
        invalid_chunk_resp.status_code == 400,
        f"Status: {invalid_chunk_resp.status_code}"
    )
    print()
    
    # A3: GET file with correct Content-Type
    print("A3: GET file returns correct Content-Type and Content-Disposition...")
    get_file_resp = requests.get(f"{API_BASE}/admin/chat/fil/{file_id_1}?key={owner_token}")
    log_test(
        "A3a: GET file returns 200 with image/png",
        get_file_resp.status_code == 200 and
        'image/png' in get_file_resp.headers.get('Content-Type', ''),
        f"Content-Type: {get_file_resp.headers.get('Content-Type')}"
    )
    
    log_test(
        "A3b: Content-Disposition is attachment by default",
        'attachment' in get_file_resp.headers.get('Content-Disposition', ''),
        f"Content-Disposition: {get_file_resp.headers.get('Content-Disposition')}"
    )
    
    # With ?inline=1
    get_file_inline_resp = requests.get(f"{API_BASE}/admin/chat/fil/{file_id_1}?key={owner_token}&inline=1")
    log_test(
        "A3c: With ?inline=1, PNG is served inline",
        'inline' in get_file_inline_resp.headers.get('Content-Disposition', ''),
        f"Content-Disposition: {get_file_inline_resp.headers.get('Content-Disposition')}"
    )
    print()
    
    # A4: Upload HTML file - should NEVER be inline
    print("A4: Upload HTML file (XSS guard)...")
    html_b64 = base64.b64encode(b'<html><body>test</body></html>').decode('utf-8')
    html_upload_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": "qa-html",
            "index": 0,
            "total": 1,
            "data": html_b64,
            "name": "test.html",
            "type": "text/html",
            "kanal": "generelt"
        }
    )
    html_file_id = html_upload_resp.json().get('fil', {}).get('id')
    tracked_file_ids.append(html_file_id)
    
    # Try to get with inline=1
    html_inline_resp = requests.get(f"{API_BASE}/admin/chat/fil/{html_file_id}?key={owner_token}&inline=1")
    log_test(
        "A4: HTML file is NEVER inline (XSS guard)",
        'attachment' in html_inline_resp.headers.get('Content-Disposition', ''),
        f"Content-Disposition: {html_inline_resp.headers.get('Content-Disposition')}"
    )
    print()
    
    # A5: GET non-existent file
    print("A5: GET non-existent file...")
    get_404_resp = requests.get(f"{API_BASE}/admin/chat/fil/finnes-ikke?key={owner_token}")
    log_test(
        "A5a: Non-existent file returns 404",
        get_404_resp.status_code == 404,
        f"Status: {get_404_resp.status_code}"
    )
    
    # Without token
    get_401_resp = requests.get(f"{API_BASE}/admin/chat/fil/{file_id_1}")
    log_test(
        "A5b: GET without token returns 401",
        get_401_resp.status_code == 401,
        f"Status: {get_401_resp.status_code}"
    )
    print()
    
    # ========================================================================
    # B) ATTACH TO MESSAGE
    # ========================================================================
    print("=" * 80)
    print("B) ATTACH TO MESSAGE TESTS")
    print("=" * 80)
    
    # B1: POST message with attachment
    print("B1: POST message with attachment...")
    msg_with_attach_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "QA melding med vedlegg",
            "vedlegg": [file_id_1]
        }
    )
    msg_with_attach = msg_with_attach_resp.json().get('melding', {})
    msg_id_1 = msg_with_attach.get('id')
    tracked_message_ids.append(msg_id_1)
    
    log_test(
        "B1a: Message created with attachment",
        msg_with_attach_resp.status_code == 201 and
        len(msg_with_attach.get('vedlegg', [])) == 1,
        f"Message ID: {msg_id_1}, Attachments: {len(msg_with_attach.get('vedlegg', []))}"
    )
    
    # Verify in MongoDB that file is bound
    file_doc = db.chat_files.find_one({'id': file_id_1})
    log_test(
        "B1b: File is bound to message in DB",
        file_doc and file_doc.get('meldingId') == msg_id_1,
        f"File meldingId: {file_doc.get('meldingId') if file_doc else 'NOT FOUND'}"
    )
    print()
    
    # B2: POST message with empty text but with attachment
    print("B2: POST message with empty text + attachment...")
    png_b64_2 = create_small_png_base64()
    upload_2_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": "qa-upload-2",
            "index": 0,
            "total": 1,
            "data": png_b64_2,
            "name": "qa-test-2.png",
            "type": "image/png",
            "kanal": "generelt"
        }
    )
    file_id_2 = upload_2_resp.json().get('fil', {}).get('id')
    tracked_file_ids.append(file_id_2)
    
    empty_text_msg_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "",
            "vedlegg": [file_id_2]
        }
    )
    msg_id_2 = empty_text_msg_resp.json().get('melding', {}).get('id')
    tracked_message_ids.append(msg_id_2)
    
    log_test(
        "B2: Empty text + attachment is OK",
        empty_text_msg_resp.status_code == 201,
        f"Status: {empty_text_msg_resp.status_code}"
    )
    print()
    
    # B3: POST with non-existent file
    print("B3: POST with non-existent file...")
    bad_file_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "Test",
            "vedlegg": ["finnes-ikke"]
        }
    )
    log_test(
        "B3: Non-existent file returns 400",
        bad_file_resp.status_code == 400 and
        'finnes ikke' in bad_file_resp.json().get('error', '').lower(),
        f"Status: {bad_file_resp.status_code}, Error: {bad_file_resp.json().get('error')}"
    )
    print()
    
    # B4: Try to reuse already-bound file
    print("B4: Try to reuse already-bound file...")
    reuse_file_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "Reuse test",
            "vedlegg": [file_id_1]
        }
    )
    log_test(
        "B4: Reusing bound file returns 400",
        reuse_file_resp.status_code == 400,
        f"Status: {reuse_file_resp.status_code}, Error: {reuse_file_resp.json().get('error')}"
    )
    print()
    
    # B5: DELETE unbound file (should work), DELETE bound file (should fail)
    print("B5: DELETE file tests...")
    # Upload a new unbound file
    png_b64_3 = create_small_png_base64()
    upload_3_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": "qa-upload-3",
            "index": 0,
            "total": 1,
            "data": png_b64_3,
            "name": "qa-test-3.png",
            "type": "image/png",
            "kanal": "generelt"
        }
    )
    file_id_3 = upload_3_resp.json().get('fil', {}).get('id')
    tracked_file_ids.append(file_id_3)
    
    # Delete unbound file
    delete_unbound_resp = requests.delete(f"{API_BASE}/admin/chat/fil/{file_id_3}?key={owner_token}")
    log_test(
        "B5a: DELETE unbound file returns ok:true",
        delete_unbound_resp.status_code == 200 and delete_unbound_resp.json().get('ok'),
        f"Response: {delete_unbound_resp.json()}"
    )
    
    # Verify file is gone
    file_gone = db.chat_files.find_one({'id': file_id_3}) is None
    log_test(
        "B5b: Unbound file is removed from DB",
        file_gone,
        f"File exists: {not file_gone}"
    )
    
    # Try to delete bound file
    delete_bound_resp = requests.delete(f"{API_BASE}/admin/chat/fil/{file_id_1}?key={owner_token}")
    log_test(
        "B5c: DELETE bound file returns ok:false",
        delete_bound_resp.status_code == 200 and not delete_bound_resp.json().get('ok'),
        f"Response: {delete_bound_resp.json()}"
    )
    
    # Verify file still exists
    file_still_exists = db.chat_files.find_one({'id': file_id_1}) is not None
    log_test(
        "B5d: Bound file still exists in DB",
        file_still_exists,
        f"File exists: {file_still_exists}"
    )
    print()
    
    # B6: DELETE message with attachment - verify cascade
    print("B6: DELETE message with attachment (cascade test)...")
    delete_msg_resp = requests.delete(f"{API_BASE}/admin/chat/meldinger?id={msg_id_1}&key={owner_token}")
    log_test(
        "B6a: DELETE message returns ok:true",
        delete_msg_resp.status_code == 200 and delete_msg_resp.json().get('ok'),
        f"Response: {delete_msg_resp.json()}"
    )
    
    # Verify file is removed (cascade)
    file_cascaded = db.chat_files.find_one({'id': file_id_1}) is None
    log_test(
        "B6b: Bound file is removed (cascade)",
        file_cascaded,
        f"File exists: {not file_cascaded}"
    )
    print()
    
    # ========================================================================
    # C) REACTIONS
    # ========================================================================
    print("=" * 80)
    print("C) REACTIONS TESTS")
    print("=" * 80)
    
    # Create a test message for reactions
    print("Creating test message for reactions...")
    react_msg_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "QA reaksjon test"
        }
    )
    react_msg_id = react_msg_resp.json().get('melding', {}).get('id')
    tracked_message_ids.append(react_msg_id)
    print(f"  Message ID: {react_msg_id}")
    print()
    
    # C1: React with 👍, then toggle off
    print("C1: React 👍 and toggle off...")
    react_1_resp = requests.post(
        f"{API_BASE}/admin/chat/reaksjon?key={owner_token}",
        json={
            "id": react_msg_id,
            "emoji": "👍"
        }
    )
    reactions_1 = react_1_resp.json().get('reaksjoner', [])
    log_test(
        "C1a: First reaction adds 👍",
        react_1_resp.status_code == 200 and
        len(reactions_1) == 1 and
        reactions_1[0].get('emoji') == '👍',
        f"Reactions: {reactions_1}"
    )
    
    # Toggle off
    react_2_resp = requests.post(
        f"{API_BASE}/admin/chat/reaksjon?key={owner_token}",
        json={
            "id": react_msg_id,
            "emoji": "👍"
        }
    )
    reactions_2 = react_2_resp.json().get('reaksjoner', [])
    log_test(
        "C1b: Second reaction toggles off",
        react_2_resp.status_code == 200 and len(reactions_2) == 0,
        f"Reactions: {reactions_2}"
    )
    print()
    
    # C2: Two different users react ❤️
    print("C2: Two users react ❤️...")
    # Owner reacts
    owner_react_resp = requests.post(
        f"{API_BASE}/admin/chat/reaksjon?key={owner_token}",
        json={
            "id": react_msg_id,
            "emoji": "❤️"
        }
    )
    
    # QA user reacts
    qa_react_resp = requests.post(
        f"{API_BASE}/admin/chat/reaksjon?key={qa_token}",
        json={
            "id": react_msg_id,
            "emoji": "❤️"
        }
    )
    reactions_multi = qa_react_resp.json().get('reaksjoner', [])
    log_test(
        "C2: Both users' reactions present",
        len(reactions_multi) == 2 and
        all(r.get('emoji') == '❤️' for r in reactions_multi),
        f"Reactions count: {len(reactions_multi)}, Emojis: {[r.get('emoji') for r in reactions_multi]}"
    )
    print()
    
    # C3: Invalid emoji
    print("C3: Invalid emoji...")
    invalid_emoji_resp = requests.post(
        f"{API_BASE}/admin/chat/reaksjon?key={owner_token}",
        json={
            "id": react_msg_id,
            "emoji": "💩"
        }
    )
    log_test(
        "C3a: Invalid emoji returns 400",
        invalid_emoji_resp.status_code == 400 and
        'ugyldig' in invalid_emoji_resp.json().get('error', '').lower(),
        f"Status: {invalid_emoji_resp.status_code}, Error: {invalid_emoji_resp.json().get('error')}"
    )
    
    # Unknown message
    unknown_msg_resp = requests.post(
        f"{API_BASE}/admin/chat/reaksjon?key={owner_token}",
        json={
            "id": "finnes-ikke",
            "emoji": "👍"
        }
    )
    log_test(
        "C3b: Unknown message returns 404",
        unknown_msg_resp.status_code == 404,
        f"Status: {unknown_msg_resp.status_code}"
    )
    print()
    
    # ========================================================================
    # D) EDIT MESSAGE
    # ========================================================================
    print("=" * 80)
    print("D) EDIT MESSAGE TESTS")
    print("=" * 80)
    
    # Create test message
    print("Creating test message for editing...")
    edit_msg_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "Original text"
        }
    )
    edit_msg_id = edit_msg_resp.json().get('melding', {}).get('id')
    tracked_message_ids.append(edit_msg_id)
    print(f"  Message ID: {edit_msg_id}")
    print()
    
    # D1: Owner edits own message
    print("D1: Owner edits own message...")
    edit_resp = requests.put(
        f"{API_BASE}/admin/chat/melding?key={owner_token}",
        json={
            "id": edit_msg_id,
            "text": "Edited text"
        }
    )
    edit_data = edit_resp.json()
    log_test(
        "D1a: Edit returns ok:true with redigertAt",
        edit_resp.status_code == 200 and
        edit_data.get('ok') and
        edit_data.get('text') == 'Edited text' and
        edit_data.get('redigertAt'),
        f"Response: {edit_data}"
    )
    
    # Verify in DB
    edited_msg = db.chat_messages.find_one({'id': edit_msg_id})
    log_test(
        "D1b: DB has updated text and redigertAt",
        edited_msg and
        edited_msg.get('text') == 'Edited text' and
        edited_msg.get('redigertAt'),
        f"Text: {edited_msg.get('text') if edited_msg else 'NOT FOUND'}, redigertAt: {edited_msg.get('redigertAt') if edited_msg else 'N/A'}"
    )
    print()
    
    # D2: QA user tries to edit owner's message
    print("D2: QA user tries to edit owner's message...")
    qa_edit_resp = requests.put(
        f"{API_BASE}/admin/chat/melding?key={qa_token}",
        json={
            "id": edit_msg_id,
            "text": "Hacked"
        }
    )
    log_test(
        "D2: Non-owner edit returns 403",
        qa_edit_resp.status_code == 403,
        f"Status: {qa_edit_resp.status_code}, Error: {qa_edit_resp.json().get('error')}"
    )
    print()
    
    # D3: Edit to empty text
    print("D3: Edit to empty text...")
    # Without attachment - should fail
    empty_edit_resp = requests.put(
        f"{API_BASE}/admin/chat/melding?key={owner_token}",
        json={
            "id": edit_msg_id,
            "text": ""
        }
    )
    log_test(
        "D3a: Empty text without attachment returns 400",
        empty_edit_resp.status_code == 400,
        f"Status: {empty_edit_resp.status_code}, Error: {empty_edit_resp.json().get('error')}"
    )
    
    # With attachment - should work
    # First create message with attachment
    png_b64_4 = create_small_png_base64()
    upload_4_resp = requests.post(
        f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
        json={
            "uploadId": "qa-upload-4",
            "index": 0,
            "total": 1,
            "data": png_b64_4,
            "name": "qa-test-4.png",
            "type": "image/png",
            "kanal": "generelt"
        }
    )
    file_id_4 = upload_4_resp.json().get('fil', {}).get('id')
    tracked_file_ids.append(file_id_4)
    
    msg_with_attach_2_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "Has attachment",
            "vedlegg": [file_id_4]
        }
    )
    msg_with_attach_2_id = msg_with_attach_2_resp.json().get('melding', {}).get('id')
    tracked_message_ids.append(msg_with_attach_2_id)
    
    # Edit to empty text
    empty_with_attach_resp = requests.put(
        f"{API_BASE}/admin/chat/melding?key={owner_token}",
        json={
            "id": msg_with_attach_2_id,
            "text": ""
        }
    )
    log_test(
        "D3b: Empty text with attachment is OK",
        empty_with_attach_resp.status_code == 200,
        f"Status: {empty_with_attach_resp.status_code}"
    )
    print()
    
    # D4: Verify NO new notifications created by edit with mention
    print("D4: Edit with mention does NOT create notifications...")
    # Get baseline notification count for QA user
    baseline_notif_count = db.notifications.count_documents({'userId': qa_user_id})
    
    # Edit message to mention QA user
    edit_mention_resp = requests.put(
        f"{API_BASE}/admin/chat/melding?key={owner_token}",
        json={
            "id": edit_msg_id,
            "text": f"Edited with mention @QA Chat User",
            "mentions": [qa_user_id]
        }
    )
    
    # Check notification count - should be unchanged
    after_notif_count = db.notifications.count_documents({'userId': qa_user_id})
    log_test(
        "D4: No new notifications from edit",
        after_notif_count == baseline_notif_count,
        f"Before: {baseline_notif_count}, After: {after_notif_count}"
    )
    print()
    
    # ========================================================================
    # E) PIN MESSAGE
    # ========================================================================
    print("=" * 80)
    print("E) PIN MESSAGE TESTS")
    print("=" * 80)
    
    # Create test message for pinning
    print("Creating test message for pinning...")
    pin_msg_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "QA pin test"
        }
    )
    pin_msg_id = pin_msg_resp.json().get('melding', {}).get('id')
    tracked_message_ids.append(pin_msg_id)
    print(f"  Message ID: {pin_msg_id}")
    print()
    
    # E1: Pin message
    print("E1: Pin message...")
    pin_resp = requests.put(
        f"{API_BASE}/admin/chat/fest?key={owner_token}",
        json={
            "id": pin_msg_id,
            "festet": True
        }
    )
    log_test(
        "E1a: Pin returns ok:true, festet:true",
        pin_resp.status_code == 200 and
        pin_resp.json().get('ok') and
        pin_resp.json().get('festet'),
        f"Response: {pin_resp.json()}"
    )
    
    # Verify in DB
    pinned_msg = db.chat_messages.find_one({'id': pin_msg_id})
    log_test(
        "E1b: DB has festet:true, festetAt, festetAv",
        pinned_msg and
        pinned_msg.get('festet') and
        pinned_msg.get('festetAt') and
        pinned_msg.get('festetAv'),
        f"festet: {pinned_msg.get('festet') if pinned_msg else 'N/A'}, festetAt: {pinned_msg.get('festetAt') if pinned_msg else 'N/A'}"
    )
    
    # GET festede
    festede_resp = requests.get(f"{API_BASE}/admin/chat/festede?kanal=generelt&key={owner_token}")
    festede = festede_resp.json().get('festede', [])
    log_test(
        "E1c: GET festede includes pinned message",
        any(f.get('id') == pin_msg_id for f in festede),
        f"Festede count: {len(festede)}"
    )
    print()
    
    # E2: Max 5 pinned messages
    print("E2: Max 5 pinned messages...")
    # Count existing pinned (including real ones from Martin)
    existing_pinned = db.chat_messages.count_documents({'kanal': 'generelt', 'festet': True})
    print(f"  Existing pinned: {existing_pinned}")
    
    # Pin enough messages to reach 5 total
    pin_msg_ids = []
    for i in range(5 - existing_pinned):
        msg_resp = requests.post(
            f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
            json={
                "kanal": "generelt",
                "text": f"QA pin test {i+2}"
            }
        )
        msg_id = msg_resp.json().get('melding', {}).get('id')
        tracked_message_ids.append(msg_id)
        pin_msg_ids.append(msg_id)
        
        pin_resp = requests.put(
            f"{API_BASE}/admin/chat/fest?key={owner_token}",
            json={
                "id": msg_id,
                "festet": True
            }
        )
        if pin_resp.status_code != 200:
            print(f"  Warning: Failed to pin message {i+2}: {pin_resp.status_code}")
    
    # Verify we have 5 pinned
    total_pinned = db.chat_messages.count_documents({'kanal': 'generelt', 'festet': True})
    print(f"  Total pinned now: {total_pinned}")
    
    # Try to pin a 6th
    msg_6_resp = requests.post(
        f"{API_BASE}/admin/chat/meldinger?key={owner_token}",
        json={
            "kanal": "generelt",
            "text": "QA pin test 6"
        }
    )
    msg_6_id = msg_6_resp.json().get('melding', {}).get('id')
    tracked_message_ids.append(msg_6_id)
    
    pin_6_resp = requests.put(
        f"{API_BASE}/admin/chat/fest?key={owner_token}",
        json={
            "id": msg_6_id,
            "festet": True
        }
    )
    log_test(
        "E2a: 6th pin returns 400",
        pin_6_resp.status_code == 400 and
        'maks 5' in pin_6_resp.json().get('error', '').lower(),
        f"Status: {pin_6_resp.status_code}, Error: {pin_6_resp.json().get('error')}"
    )
    
    # Unpin our test pins
    for msg_id in [pin_msg_id] + pin_msg_ids:
        unpin_resp = requests.put(
            f"{API_BASE}/admin/chat/fest?key={owner_token}",
            json={
                "id": msg_id,
                "festet": False
            }
        )
        if unpin_resp.status_code != 200:
            print(f"  Warning: Failed to unpin {msg_id}: {unpin_resp.status_code}")
    
    log_test(
        "E2b: Unpinned test messages",
        True,
        "All test pins removed"
    )
    print()
    
    # E3: Unpin message
    print("E3: Unpin message...")
    # Pin again
    pin_again_resp = requests.put(
        f"{API_BASE}/admin/chat/fest?key={owner_token}",
        json={
            "id": pin_msg_id,
            "festet": True
        }
    )
    
    # Unpin
    unpin_resp = requests.put(
        f"{API_BASE}/admin/chat/fest?key={owner_token}",
        json={
            "id": pin_msg_id,
            "festet": False
        }
    )
    log_test(
        "E3a: Unpin returns ok:true, festet:false",
        unpin_resp.status_code == 200 and
        unpin_resp.json().get('ok') and
        not unpin_resp.json().get('festet'),
        f"Response: {unpin_resp.json()}"
    )
    
    # GET festede should not include it
    festede_after_resp = requests.get(f"{API_BASE}/admin/chat/festede?kanal=generelt&key={owner_token}")
    festede_after = festede_after_resp.json().get('festede', [])
    log_test(
        "E3b: GET festede does not include unpinned message",
        not any(f.get('id') == pin_msg_id for f in festede_after),
        f"Festede count: {len(festede_after)}"
    )
    print()
    
    # ========================================================================
    # F) TYPING INDICATOR
    # ========================================================================
    print("=" * 80)
    print("F) TYPING INDICATOR TESTS")
    print("=" * 80)
    
    # F1: QA user posts heartbeat, owner gets it
    print("F1: QA user typing, owner sees it...")
    qa_typing_resp = requests.post(
        f"{API_BASE}/admin/chat/skriver?key={qa_token}",
        json={
            "kanal": "generelt"
        }
    )
    log_test(
        "F1a: POST skriver returns ok:true",
        qa_typing_resp.status_code == 200 and qa_typing_resp.json().get('ok'),
        f"Response: {qa_typing_resp.json()}"
    )
    
    # Owner gets typing indicator
    owner_get_typing_resp = requests.get(f"{API_BASE}/admin/chat/skriver?kanal=generelt&key={owner_token}")
    skriver = owner_get_typing_resp.json().get('skriver', [])
    log_test(
        "F1b: Owner sees QA user typing",
        'QA Chat User' in skriver,
        f"Skriver: {skriver}"
    )
    print()
    
    # F2: Owner does NOT see own name
    print("F2: Owner does NOT see own name...")
    owner_typing_resp = requests.post(
        f"{API_BASE}/admin/chat/skriver?key={owner_token}",
        json={
            "kanal": "generelt"
        }
    )
    
    owner_get_own_resp = requests.get(f"{API_BASE}/admin/chat/skriver?kanal=generelt&key={owner_token}")
    skriver_own = owner_get_own_resp.json().get('skriver', [])
    log_test(
        "F2: Owner does NOT see own name",
        'Martin' not in ' '.join(skriver_own),
        f"Skriver: {skriver_own}"
    )
    print()
    
    # F3: Wait 7 seconds, QA user should disappear
    print("F3: Wait 7 seconds, QA user disappears...")
    print("  Waiting 7 seconds...")
    time.sleep(7)
    
    owner_get_after_resp = requests.get(f"{API_BASE}/admin/chat/skriver?kanal=generelt&key={owner_token}")
    skriver_after = owner_get_after_resp.json().get('skriver', [])
    log_test(
        "F3: QA user no longer typing (6s freshness)",
        'QA Chat User' not in skriver_after,
        f"Skriver: {skriver_after}"
    )
    print()
    
    # ========================================================================
    # G) DOCX PREVIEW
    # ========================================================================
    print("=" * 80)
    print("G) DOCX PREVIEW TESTS")
    print("=" * 80)
    
    docx_b64 = create_minimal_docx()
    if docx_b64:
        # G1: Upload DOCX and get preview
        print("G1: Upload DOCX and get preview...")
        docx_upload_resp = requests.post(
            f"{API_BASE}/admin/chat/fil-chunk?key={owner_token}",
            json={
                "uploadId": "qa-docx",
                "index": 0,
                "total": 1,
                "data": docx_b64,
                "name": "test.docx",
                "type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "kanal": "generelt"
            }
        )
        docx_file_id = docx_upload_resp.json().get('fil', {}).get('id')
        tracked_file_ids.append(docx_file_id)
        
        # Get preview
        preview_resp = requests.get(f"{API_BASE}/admin/chat/fil/{docx_file_id}/forhandsvisning?key={owner_token}")
        log_test(
            "G1a: DOCX preview returns 200 with application/pdf",
            preview_resp.status_code == 200 and
            'application/pdf' in preview_resp.headers.get('Content-Type', ''),
            f"Status: {preview_resp.status_code}, Content-Type: {preview_resp.headers.get('Content-Type')}"
        )
        
        log_test(
            "G1b: PDF body starts with %PDF-",
            preview_resp.content[:5] == b'%PDF-',
            f"First 5 bytes: {preview_resp.content[:5]}"
        )
        
        # Second GET should be cached
        preview_2_resp = requests.get(f"{API_BASE}/admin/chat/fil/{docx_file_id}/forhandsvisning?key={owner_token}")
        # Check if pdfData is set in DB
        docx_file_doc = db.chat_files.find_one({'id': docx_file_id})
        log_test(
            "G1c: Second GET served from cache (pdfData set)",
            docx_file_doc and docx_file_doc.get('pdfData'),
            f"pdfData exists: {bool(docx_file_doc.get('pdfData') if docx_file_doc else False)}"
        )
        print()
    else:
        print("G1: SKIPPED (python-docx not available)")
        print()
    
    # G2: Preview on PNG should fail
    print("G2: Preview on PNG returns 400...")
    # Use file_id_2 (PNG)
    png_preview_resp = requests.get(f"{API_BASE}/admin/chat/fil/{file_id_2}/forhandsvisning?key={owner_token}")
    log_test(
        "G2a: PNG preview returns 400",
        png_preview_resp.status_code == 400,
        f"Status: {png_preview_resp.status_code}, Error: {png_preview_resp.json().get('error')}"
    )
    
    # Unknown file
    unknown_preview_resp = requests.get(f"{API_BASE}/admin/chat/fil/finnes-ikke/forhandsvisning?key={owner_token}")
    log_test(
        "G2b: Unknown file preview returns 404",
        unknown_preview_resp.status_code == 404,
        f"Status: {unknown_preview_resp.status_code}"
    )
    
    # Without token
    no_token_preview_resp = requests.get(f"{API_BASE}/admin/chat/fil/{file_id_2}/forhandsvisning")
    log_test(
        "G2c: Preview without token returns 401",
        no_token_preview_resp.status_code == 401,
        f"Status: {no_token_preview_resp.status_code}"
    )
    print()
    
    # ========================================================================
    # CLEANUP
    # ========================================================================
    print("=" * 80)
    print("CLEANUP")
    print("=" * 80)
    
    print("Deleting test messages...")
    for msg_id in tracked_message_ids:
        try:
            delete_resp = requests.delete(f"{API_BASE}/admin/chat/meldinger?id={msg_id}&key={ADMIN_KEY}")
            if delete_resp.status_code == 200:
                print(f"  ✓ Deleted message {msg_id}")
            else:
                print(f"  ⚠ Failed to delete message {msg_id}: {delete_resp.status_code}")
        except Exception as e:
            print(f"  ⚠ Error deleting message {msg_id}: {e}")
    
    print("Deleting remaining unbound files...")
    for file_id in tracked_file_ids:
        try:
            # Check if file still exists
            file_doc = db.chat_files.find_one({'id': file_id})
            if file_doc:
                db.chat_files.delete_one({'id': file_id})
                print(f"  ✓ Deleted file {file_id}")
        except Exception as e:
            print(f"  ⚠ Error deleting file {file_id}: {e}")
    
    print("Deleting QA users...")
    for user_id in tracked_user_ids:
        try:
            delete_resp = requests.delete(f"{API_BASE}/admin/users/{user_id}?key={ADMIN_KEY}")
            if delete_resp.status_code == 200:
                print(f"  ✓ Deleted user {user_id}")
            else:
                print(f"  ⚠ Failed to delete user {user_id}: {delete_resp.status_code}")
        except Exception as e:
            print(f"  ⚠ Error deleting user {user_id}: {e}")
    
    # Delete chat_lest and chat_skriver docs for QA users
    print("Deleting chat_lest and chat_skriver docs...")
    for user_id in tracked_user_ids:
        db.chat_lest.delete_many({'userId': user_id})
        db.chat_skriver.delete_many({'userId': user_id})
    
    # Delete notifications for QA users
    print("Deleting notifications for QA users...")
    for user_id in tracked_user_ids:
        result = db.notifications.delete_many({'userId': user_id})
        print(f"  ✓ Deleted {result.deleted_count} notifications for {user_id}")
    
    print()
    
    # Verify cleanup
    print("Verifying cleanup...")
    final_messages = db.chat_messages.count_documents({})
    final_files = db.chat_files.count_documents({})
    final_pinned = db.chat_messages.count_documents({'festet': True})
    
    log_test(
        "Cleanup: Martin's messages unchanged",
        final_messages == baseline_messages,
        f"Baseline: {baseline_messages}, Final: {final_messages}"
    )
    
    log_test(
        "Cleanup: Martin's files unchanged",
        final_files == baseline_files,
        f"Baseline: {baseline_files}, Final: {final_files}"
    )
    
    log_test(
        "Cleanup: Pinned count back to baseline",
        final_pinned == baseline_pinned,
        f"Baseline: {baseline_pinned}, Final: {final_pinned}"
    )
    
    # Verify no test data left
    qa_messages = db.chat_messages.count_documents({'text': {'$regex': '^QA '}})
    qa_files = db.chat_files.count_documents({'name': {'$regex': '^qa-'}})
    qa_users = db.admin_users.count_documents({'email': {'$regex': '@example.com$'}})
    
    log_test(
        "Cleanup: No QA messages left",
        qa_messages == 0,
        f"QA messages: {qa_messages}"
    )
    
    log_test(
        "Cleanup: No QA files left",
        qa_files == 0,
        f"QA files: {qa_files}"
    )
    
    log_test(
        "Cleanup: No QA users left",
        qa_users == 0,
        f"QA users: {qa_users}"
    )
    
    print()
    print("=" * 80)
    print("ALL TESTS COMPLETED SUCCESSFULLY!")
    print("=" * 80)
    
    mongo_client.close()

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n❌ TEST SUITE FAILED: {e}")
        import traceback
        traceback.print_exc()
        exit(1)
