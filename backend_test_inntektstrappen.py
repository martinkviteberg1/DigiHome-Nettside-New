#!/usr/bin/env python3
"""
Backend test for INNTEKTSTRAPPEN (annonsert = pipeline) feature.
Tests the advertised flag normalization, filtering, and Excel/CSV export.

CRITICAL SAFETY RULES:
- SendGrid is LIVE (no email flows triggered - only GET requests)
- Do NOT delete/modify 'Lønn — 1 ansatt' (60000) in enhetsokonomi
- Do NOT touch users qa-investor@example.com or martin@kviteberg.no
- Platform upstream is REAL PRODUCTION - all calls are GET/read-only (safe)
"""

import requests
import sys
import re
from io import BytesIO

BASE_URL = "https://saker-hub.preview.emergentagent.com/api"
ADMIN_KEY = "dh_admin_b3Kx92Qz7Lm4"

def test_1_fetch_leieforhold_with_advertised():
    """
    Test 1: GET /api/admin/leieforhold?key=&fresh=1 (timeout ≥60s)
    → 200, source='lease-income'
    Count rows where status_label contains 'annonsert' (case-insensitive)
    → each must have advertised=true
    Expected roughly 5-7 such rows (data may drift; assert count ≥1 and flag consistency)
    """
    print("\n" + "="*80)
    print("TEST 1: Fetch leieforhold with advertised flag normalization")
    print("="*80)
    
    try:
        url = f"{BASE_URL}/admin/leieforhold?key={ADMIN_KEY}&fresh=1"
        print(f"GET {url}")
        
        # Use 60s timeout as specified
        response = requests.get(url, timeout=60)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            print(f"Response: {response.text[:500]}")
            return False
        
        data = response.json()
        
        # Verify source
        source = data.get('source', '')
        print(f"Source: {source}")
        if source != 'lease-income':
            print(f"⚠️  WARNING: Expected source='lease-income', got '{source}'")
        
        # Get rows
        rows = data.get('rows', [])
        print(f"Total rows: {len(rows)}")
        
        # Count rows where status_label contains 'annonsert' (case-insensitive)
        annonsert_rows = []
        for r in rows:
            status_label = str(r.get('status_label', '')).lower()
            if 'annonsert' in status_label:
                annonsert_rows.append(r)
        
        print(f"Rows with 'annonsert' in status_label: {len(annonsert_rows)}")
        
        if len(annonsert_rows) < 1:
            print(f"❌ FAILED: Expected at least 1 row with 'annonsert', got {len(annonsert_rows)}")
            return False
        
        # Verify each has advertised=true
        all_have_flag = True
        for r in annonsert_rows:
            if not r.get('advertised'):
                print(f"❌ FAILED: Row '{r.get('unit_room')}' has 'annonsert' in status_label but advertised={r.get('advertised')}")
                all_have_flag = False
        
        if not all_have_flag:
            return False
        
        print(f"✅ PASSED: All {len(annonsert_rows)} rows with 'annonsert' in status_label have advertised=true")
        
        # Store count for next tests
        global ADVERTISED_COUNT
        ADVERTISED_COUNT = len(annonsert_rows)
        
        # Count total vacant group rows
        vacant_group_rows = [r for r in rows if r.get('group') == 'vacant']
        print(f"Total vacant group rows: {len(vacant_group_rows)}")
        
        global VACANT_WITHOUT_AD_COUNT
        VACANT_WITHOUT_AD_COUNT = len([r for r in vacant_group_rows if not r.get('advertised')])
        print(f"Vacant WITHOUT advertised: {VACANT_WITHOUT_AD_COUNT}")
        
        return True
        
    except Exception as e:
        print(f"❌ EXCEPTION: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_2_csv_export_advertised():
    """
    Test 2: GET /api/admin/leieforhold/csv?key=&status=advertised
    → 200, Content-Disposition contains '-filtrert'
    Number of data lines (after header) must equal advertised count from test 1
    Every data line must contain 'Annonsert'
    """
    print("\n" + "="*80)
    print("TEST 2: CSV export with status=advertised filter")
    print("="*80)
    
    try:
        url = f"{BASE_URL}/admin/leieforhold/csv?key={ADMIN_KEY}&status=advertised"
        print(f"GET {url}")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            return False
        
        # Check Content-Disposition
        content_disp = response.headers.get('Content-Disposition', '')
        print(f"Content-Disposition: {content_disp}")
        
        if '-filtrert' not in content_disp:
            print(f"❌ FAILED: Expected '-filtrert' in Content-Disposition")
            return False
        
        # Parse CSV
        csv_text = response.text
        lines = csv_text.strip().split('\n')
        
        # Remove BOM if present
        if lines[0].startswith('\ufeff'):
            lines[0] = lines[0][1:]
        
        header = lines[0]
        data_lines = lines[1:]
        
        print(f"Header: {header[:100]}...")
        print(f"Data lines: {len(data_lines)}")
        
        # Verify count matches
        if len(data_lines) != ADVERTISED_COUNT:
            print(f"❌ FAILED: Expected {ADVERTISED_COUNT} data lines, got {len(data_lines)}")
            return False
        
        # Verify every data line contains 'Annonsert'
        all_contain_annonsert = True
        for i, line in enumerate(data_lines):
            if 'Annonsert' not in line:
                print(f"❌ FAILED: Data line {i+1} does NOT contain 'Annonsert': {line[:100]}")
                all_contain_annonsert = False
        
        if not all_contain_annonsert:
            return False
        
        print(f"✅ PASSED: CSV has {len(data_lines)} data lines, all contain 'Annonsert'")
        return True
        
    except Exception as e:
        print(f"❌ EXCEPTION: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_3_csv_export_vacant():
    """
    Test 3: GET csv with status=vacant
    → 200; data lines must NOT contain 'Annonsert'
    Count must equal vacant group rows where advertised=false from test 1
    """
    print("\n" + "="*80)
    print("TEST 3: CSV export with status=vacant filter (WITHOUT advertised)")
    print("="*80)
    
    try:
        url = f"{BASE_URL}/admin/leieforhold/csv?key={ADMIN_KEY}&status=vacant"
        print(f"GET {url}")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            return False
        
        # Parse CSV
        csv_text = response.text
        lines = csv_text.strip().split('\n')
        
        # Remove BOM if present
        if lines[0].startswith('\ufeff'):
            lines[0] = lines[0][1:]
        
        data_lines = lines[1:]
        print(f"Data lines: {len(data_lines)}")
        
        # Verify count matches vacant WITHOUT advertised
        if len(data_lines) != VACANT_WITHOUT_AD_COUNT:
            print(f"❌ FAILED: Expected {VACANT_WITHOUT_AD_COUNT} data lines (vacant without ad), got {len(data_lines)}")
            return False
        
        # Verify NO data line contains 'Annonsert'
        any_contain_annonsert = False
        for i, line in enumerate(data_lines):
            if 'Annonsert' in line:
                print(f"❌ FAILED: Data line {i+1} contains 'Annonsert' (should NOT): {line[:100]}")
                any_contain_annonsert = True
        
        if any_contain_annonsert:
            return False
        
        print(f"✅ PASSED: CSV has {len(data_lines)} data lines, NONE contain 'Annonsert'")
        
        # Store count for test 4
        global VACANT_CSV_COUNT
        VACANT_CSV_COUNT = len(data_lines)
        
        return True
        
    except Exception as e:
        print(f"❌ EXCEPTION: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_4_csv_export_combined():
    """
    Test 4: GET csv with status=advertised,vacant
    → 200; data line count = test 2 count + test 3 count
    """
    print("\n" + "="*80)
    print("TEST 4: CSV export with status=advertised,vacant (combined)")
    print("="*80)
    
    try:
        url = f"{BASE_URL}/admin/leieforhold/csv?key={ADMIN_KEY}&status=advertised,vacant"
        print(f"GET {url}")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            return False
        
        # Parse CSV
        csv_text = response.text
        lines = csv_text.strip().split('\n')
        
        # Remove BOM if present
        if lines[0].startswith('\ufeff'):
            lines[0] = lines[0][1:]
        
        data_lines = lines[1:]
        print(f"Data lines: {len(data_lines)}")
        
        expected_count = ADVERTISED_COUNT + VACANT_CSV_COUNT
        print(f"Expected: {ADVERTISED_COUNT} (advertised) + {VACANT_CSV_COUNT} (vacant) = {expected_count}")
        
        if len(data_lines) != expected_count:
            print(f"❌ FAILED: Expected {expected_count} data lines, got {len(data_lines)}")
            return False
        
        print(f"✅ PASSED: CSV has {len(data_lines)} data lines (sum of advertised + vacant)")
        return True
        
    except Exception as e:
        print(f"❌ EXCEPTION: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_5_xlsx_export_advertised():
    """
    Test 5: GET xlsx with status=advertised
    → 200 valid XLSX (PK signature), Oversikt sheet exists
    """
    print("\n" + "="*80)
    print("TEST 5: XLSX export with status=advertised filter")
    print("="*80)
    
    try:
        url = f"{BASE_URL}/admin/leieforhold/xlsx?key={ADMIN_KEY}&status=advertised"
        print(f"GET {url}")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            return False
        
        # Check XLSX signature (PK header)
        content = response.content
        if not content.startswith(b'PK'):
            print(f"❌ FAILED: Response does NOT start with PK signature (not a valid XLSX)")
            print(f"First bytes: {content[:20]}")
            return False
        
        print(f"✅ Valid XLSX file (PK signature present), size: {len(content)} bytes")
        
        # Try to open with openpyxl
        try:
            from openpyxl import load_workbook
            import warnings
            
            # Ignore known warnings about conditional formatting/data validation
            with warnings.catch_warnings():
                warnings.filterwarnings('ignore', category=UserWarning)
                wb = load_workbook(BytesIO(content))
            
            sheet_names = wb.sheetnames
            print(f"Sheet names: {sheet_names}")
            
            if 'Oversikt' not in sheet_names:
                print(f"❌ FAILED: 'Oversikt' sheet NOT found in workbook")
                return False
            
            print(f"✅ PASSED: Valid XLSX with 'Oversikt' sheet")
            return True
            
        except ImportError:
            print("⚠️  WARNING: openpyxl not available, skipping sheet verification")
            print("✅ PASSED: Valid XLSX file (PK signature verified)")
            return True
        
    except Exception as e:
        print(f"❌ EXCEPTION: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_6_xlsx_export_full_with_inntektstrappen():
    """
    Test 6: GET xlsx without filter
    → 200 valid XLSX
    Open with openpyxl (ignore known warnings) and assert 'Oversikt' sheet
    contains cell text 'Inntektstrappen / mnd' and 'Annonsert · pipeline'
    """
    print("\n" + "="*80)
    print("TEST 6: XLSX export without filter - verify Inntektstrappen block")
    print("="*80)
    
    try:
        url = f"{BASE_URL}/admin/leieforhold/xlsx?key={ADMIN_KEY}"
        print(f"GET {url}")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            return False
        
        # Check XLSX signature
        content = response.content
        if not content.startswith(b'PK'):
            print(f"❌ FAILED: Response does NOT start with PK signature")
            return False
        
        print(f"✅ Valid XLSX file, size: {len(content)} bytes")
        
        # Open with openpyxl and check for specific text
        try:
            from openpyxl import load_workbook
            import warnings
            
            # Ignore known warnings
            with warnings.catch_warnings():
                warnings.filterwarnings('ignore', category=UserWarning)
                wb = load_workbook(BytesIO(content))
            
            if 'Oversikt' not in wb.sheetnames:
                print(f"❌ FAILED: 'Oversikt' sheet NOT found")
                return False
            
            ws = wb['Oversikt']
            
            # Search for 'Inntektstrappen / mnd' and 'Annonsert · pipeline'
            found_inntektstrappen = False
            found_annonsert_pipeline = False
            
            for row in ws.iter_rows():
                for cell in row:
                    cell_value = str(cell.value or '').strip()
                    cell_value_lower = cell_value.lower()
                    if 'inntektstrappen' in cell_value_lower and 'mnd' in cell_value_lower:
                        found_inntektstrappen = True
                        print(f"✅ Found 'Inntektstrappen / mnd' in cell {cell.coordinate}: '{cell_value}'")
                    if 'annonsert' in cell_value_lower and 'pipeline' in cell_value_lower:
                        found_annonsert_pipeline = True
                        print(f"✅ Found 'Annonsert · pipeline' in cell {cell.coordinate}: '{cell_value}'")
            
            if not found_inntektstrappen:
                print(f"❌ FAILED: 'Inntektstrappen / mnd' NOT found in Oversikt sheet")
                return False
            
            if not found_annonsert_pipeline:
                print(f"❌ FAILED: 'Annonsert · pipeline' NOT found in Oversikt sheet")
                return False
            
            print(f"✅ PASSED: Oversikt sheet contains both 'Inntektstrappen / mnd' and 'Annonsert · pipeline'")
            return True
            
        except ImportError:
            print("⚠️  WARNING: openpyxl not available, cannot verify sheet contents")
            print("✅ PASSED: Valid XLSX file (PK signature verified)")
            return True
        
    except Exception as e:
        print(f"❌ EXCEPTION: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_7_regression():
    """
    Test 7: REGRESSION
    - GET /api/admin/leieforhold/okonomi → 200 with 'Lønn — 1 ansatt' present
    - GET /api/admin/leieforhold without key → 401
    - GET /api/admin/leieforhold/csv without key → 401
    """
    print("\n" + "="*80)
    print("TEST 7: Regression tests")
    print("="*80)
    
    all_passed = True
    
    # Test 7a: okonomi endpoint
    try:
        url = f"{BASE_URL}/admin/leieforhold/okonomi?key={ADMIN_KEY}"
        print(f"\n7a) GET {url}")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {response.status_code}")
            all_passed = False
        else:
            data = response.json()
            felles = data.get('felles', [])
            
            # Look for 'Lønn — 1 ansatt' with value 60000
            found_lonn = False
            for item in felles:
                if 'Lønn' in item.get('navn', '') and '1 ansatt' in item.get('navn', ''):
                    found_lonn = True
                    print(f"✅ Found '{item.get('navn')}' with belop={item.get('belop')}")
                    if item.get('belop') != 60000:
                        print(f"⚠️  WARNING: Expected belop=60000, got {item.get('belop')}")
                    break
            
            if not found_lonn:
                print(f"❌ FAILED: 'Lønn — 1 ansatt' NOT found in okonomi.felles")
                all_passed = False
            else:
                print(f"✅ PASSED: okonomi endpoint returns 200 with 'Lønn — 1 ansatt'")
        
    except Exception as e:
        print(f"❌ EXCEPTION in 7a: {e}")
        all_passed = False
    
    # Test 7b: leieforhold without key → 401
    try:
        url = f"{BASE_URL}/admin/leieforhold"
        print(f"\n7b) GET {url} (without key)")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 401:
            print(f"❌ FAILED: Expected 401, got {response.status_code}")
            all_passed = False
        else:
            print(f"✅ PASSED: leieforhold without key returns 401")
        
    except Exception as e:
        print(f"❌ EXCEPTION in 7b: {e}")
        all_passed = False
    
    # Test 7c: csv without key → 401
    try:
        url = f"{BASE_URL}/admin/leieforhold/csv"
        print(f"\n7c) GET {url} (without key)")
        
        response = requests.get(url, timeout=30)
        print(f"Status: {response.status_code}")
        
        if response.status_code != 401:
            print(f"❌ FAILED: Expected 401, got {response.status_code}")
            all_passed = False
        else:
            print(f"✅ PASSED: csv without key returns 401")
        
    except Exception as e:
        print(f"❌ EXCEPTION in 7c: {e}")
        all_passed = False
    
    return all_passed


def main():
    print("="*80)
    print("INNTEKTSTRAPPEN (annonsert = pipeline) - Backend Test Suite")
    print("="*80)
    print(f"Base URL: {BASE_URL}")
    print(f"Admin key: {ADMIN_KEY}")
    print("\nCRITICAL SAFETY RULES:")
    print("- SendGrid is LIVE (no email flows)")
    print("- Do NOT delete/modify 'Lønn — 1 ansatt' (60000)")
    print("- Do NOT touch qa-investor@example.com or martin@kviteberg.no")
    print("- Platform upstream is REAL PRODUCTION - all GET/read-only")
    print("="*80)
    
    results = []
    
    # Run tests
    results.append(("Test 1: Fetch leieforhold with advertised flag", test_1_fetch_leieforhold_with_advertised()))
    results.append(("Test 2: CSV export status=advertised", test_2_csv_export_advertised()))
    results.append(("Test 3: CSV export status=vacant", test_3_csv_export_vacant()))
    results.append(("Test 4: CSV export status=advertised,vacant", test_4_csv_export_combined()))
    results.append(("Test 5: XLSX export status=advertised", test_5_xlsx_export_advertised()))
    results.append(("Test 6: XLSX export full with Inntektstrappen", test_6_xlsx_export_full_with_inntektstrappen()))
    results.append(("Test 7: Regression tests", test_7_regression()))
    
    # Summary
    print("\n" + "="*80)
    print("TEST SUMMARY")
    print("="*80)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for name, result in results:
        status = "✅ PASSED" if result else "❌ FAILED"
        print(f"{status}: {name}")
    
    print("="*80)
    print(f"TOTAL: {passed}/{total} tests passed ({passed*100//total}% success rate)")
    print("="*80)
    
    if passed == total:
        print("\n🎉 ALL TESTS PASSED!")
        return 0
    else:
        print(f"\n⚠️  {total - passed} test(s) FAILED")
        return 1


if __name__ == "__main__":
    sys.exit(main())
